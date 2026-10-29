"""
resume_parser.py

Simple, offline resume parser -> plain text.
Supports: .pdf (pdfplumber), .docx (python-docx)
Handles:  English and Persian (RTL, presentation forms, digits, ZWNJ)

Usage (CLI):
    python resume_parser.py path/to/resume.pdf
    python resume_parser.py path/to/resume.docx --out resume.txt

Usage (library):
    from resume_parser import parse_resume
    text = parse_resume(open("cv.pdf", "rb").read(), "cv.pdf")
"""

from __future__ import annotations

import io
import re
import sys
import argparse
import unicodedata
from pathlib import Path

try:
    import pdfplumber
except ImportError:
    pdfplumber = None

try:
    import docx  # python-docx
except ImportError:
    docx = None

try:
    from bidi.algorithm import get_display  # python-bidi
except ImportError:
    get_display = None


# =============================================================================
# PERSIAN / ARABIC NORMALIZATION
# =============================================================================

# Arabic presentation forms -> base letters.
# PDF extractors often leave these in, which turns Persian into gibberish
# for any downstream text processing. We map the common ranges back.
_PRESENTATION_MAP = {
    **{chr(c): "\u0627" for c in range(0xFE8D, 0xFE8F)},   # ALEF
    **{chr(c): "\u0628" for c in range(0xFE8F, 0xFE93)},   # BEH
    **{chr(c): "\u067E" for c in range(0xFB56, 0xFB5A)},   # PEH
    **{chr(c): "\u062A" for c in range(0xFE95, 0xFE99)},   # TEH
    **{chr(c): "\u062B" for c in range(0xFE99, 0xFE9D)},   # THEH
    **{chr(c): "\u062C" for c in range(0xFE9D, 0xFEA1)},   # JEEM
    **{chr(c): "\u0686" for c in range(0xFB7A, 0xFB7E)},   # TCHEH
    **{chr(c): "\u062D" for c in range(0xFEA1, 0xFEA5)},   # HAH
    **{chr(c): "\u062E" for c in range(0xFEA5, 0xFEA9)},   # KHAH
    **{chr(c): "\u062F" for c in range(0xFEA9, 0xFEAB)},   # DAL
    **{chr(c): "\u0630" for c in range(0xFEAB, 0xFEAD)},   # THAL
    **{chr(c): "\u0631" for c in range(0xFEAD, 0xFEAF)},   # REH
    **{chr(c): "\u0632" for c in range(0xFEAF, 0xFEB1)},   # ZAIN
    **{chr(c): "\u0698" for c in range(0xFB8A, 0xFB8C)},   # JEH
    **{chr(c): "\u0633" for c in range(0xFEB1, 0xFEB5)},   # SEEN
    **{chr(c): "\u0634" for c in range(0xFEB5, 0xFEB9)},   # SHEEN
    **{chr(c): "\u0635" for c in range(0xFEB9, 0xFEBD)},   # SAD
    **{chr(c): "\u0636" for c in range(0xFEBD, 0xFEC1)},   # DAD
    **{chr(c): "\u0637" for c in range(0xFEC1, 0xFEC5)},   # TAH
    **{chr(c): "\u0638" for c in range(0xFEC5, 0xFEC9)},   # ZAH
    **{chr(c): "\u0639" for c in range(0xFEC9, 0xFECD)},   # AIN
    **{chr(c): "\u063A" for c in range(0xFECD, 0xFED1)},   # GHAIN
    **{chr(c): "\u0641" for c in range(0xFED1, 0xFED5)},   # FEH
    **{chr(c): "\u0642" for c in range(0xFED5, 0xFED9)},   # QAF
    **{chr(c): "\u06A9" for c in range(0xFB8E, 0xFB92)},   # KEHEH
    **{chr(c): "\u06AF" for c in range(0xFB92, 0xFB96)},   # GAF
    **{chr(c): "\u0644" for c in range(0xFEDD, 0xFEE1)},   # LAM
    **{chr(c): "\u0645" for c in range(0xFEE1, 0xFEE5)},   # MEEM
    **{chr(c): "\u0646" for c in range(0xFEE5, 0xFEE9)},   # NOON
    **{chr(c): "\u0647" for c in range(0xFEE9, 0xFEED)},   # HEH
    **{chr(c): "\u0648" for c in range(0xFEED, 0xFEEF)},   # WAW
    **{chr(c): "\u06CC" for c in range(0xFBFC, 0xFC00)},   # FARSI YEH
    **{chr(c): "\u06CC" for c in range(0xFEF1, 0xFEF5)},   # YEH
}

# Arabic-Indic + Persian digits -> ASCII digits
_DIGIT_MAP = str.maketrans(
    "\u0660\u0661\u0662\u0663\u0664\u0665\u0666\u0667\u0668\u0669"
    "\u06F0\u06F1\u06F2\u06F3\u06F4\u06F5\u06F6\u06F7\u06F8\u06F9",
    "01234567890123456789",
)

# Persian punctuation -> ASCII (helps downstream text matching)
_PUNCT_MAP = str.maketrans({
    "\u060C": ",", "\u061B": ";", "\u061F": "?", "\u066A": "%",
    "\u066B": ".", "\u066C": ",", "\u2010": "-", "\u2011": "-",
    "\u2012": "-", "\u2013": "-", "\u2014": "-",
    "\u2018": "'", "\u2019": "'", "\u201C": '"', "\u201D": '"',
    "\u00AB": '"', "\u00BB": '"',
    "\u200E": "", "\u200F": "",   # LRM / RLM
})

# Arabic letter variants -> Persian equivalents (unifies lookalike letters)
_LETTER_MAP = str.maketrans({
    "\u064A": "\u06CC",   # ي -> ی
    "\u0649": "\u06CC",   # ى -> ی
    "\u0643": "\u06A9",   # ك -> ک
    "\u0629": "\u0647",   # ة -> ه
})


def _strip_presentation_forms(text: str) -> str:
    return "".join(_PRESENTATION_MAP.get(ch, ch) for ch in text)


def _has_presentation_forms(text: str) -> bool:
    return any("\uFB50" <= ch <= "\uFEFF" for ch in text)


def normalize_text(text: str) -> str:
    """Full Persian/English normalization. Safe for pure English input too."""
    if not text:
        return ""
    text = _strip_presentation_forms(text)
    text = unicodedata.normalize("NFC", text)
    text = text.translate(_LETTER_MAP)
    text = text.translate(_DIGIT_MAP)
    text = text.translate(_PUNCT_MAP)
    text = text.replace("\u200C", " ").replace("\u200D", "")   # ZWNJ/ZWJ
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


# =============================================================================
# PARSERS
# =============================================================================

def _parse_pdf(file_bytes: bytes) -> str:
    if pdfplumber is None:
        raise RuntimeError("pdfplumber not installed. pip install pdfplumber")

    parts: list[str] = []
    with pdfplumber.open(io.BytesIO(file_bytes)) as pdf:
        for page in pdf.pages:
            parts.append(page.extract_text() or "")
    raw = "\n".join(parts)

    # If glyphs are visual-order Persian, run bidi to recover logical order.
    if _has_presentation_forms(raw) and get_display is not None:
        raw = normalize_text(raw)
        raw = "\n".join(get_display(line) for line in raw.splitlines())

    return raw


def _parse_docx(file_bytes: bytes) -> str:
    if docx is None:
        raise RuntimeError("python-docx not installed. pip install python-docx")

    document = docx.Document(io.BytesIO(file_bytes))
    parts: list[str] = []

    # Paragraphs
    for p in document.paragraphs:
        if p.text:
            parts.append(p.text)

    # Tables (skills are often in tables, especially on Persian resumes)
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text and c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))

    # Headers / footers
    for section in document.sections:
        for hf in (section.header, section.footer):
            for p in hf.paragraphs:
                if p.text.strip():
                    parts.append(p.text)

    return "\n".join(parts)


# =============================================================================
# PUBLIC API
# =============================================================================

def parse_resume(file_bytes: bytes, filename: str) -> str:
    """Extract plain text from a PDF or DOCX resume (Persian or English)."""
    lower = filename.lower()
    if lower.endswith(".pdf"):
        raw = _parse_pdf(file_bytes)
    elif lower.endswith(".docx"):
        raw = _parse_docx(file_bytes)
    else:
        raise ValueError(f"Unsupported file: {filename} (use .pdf or .docx)")
    return normalize_text(raw)


def parse_resume_file(path: str | Path) -> str:
    p = Path(path)
    return parse_resume(p.read_bytes(), p.name)


# =============================================================================
# CLI
# =============================================================================

def _cli() -> None:
    ap = argparse.ArgumentParser(description="Parse a PDF/DOCX resume to plain text.")
    ap.add_argument("path", help="Path to .pdf or .docx")
    ap.add_argument("--out", help="Write to this .txt file (default: stdout)")
    args = ap.parse_args()

    p = Path(args.path)
    if not p.exists():
        sys.exit(f"File not found: {p}")

    text = parse_resume_file(p)
    if args.out:
        Path(args.out).write_text(text, encoding="utf-8")
        print(f"Wrote {len(text)} characters to {args.out}")
    else:
        print(text)


if __name__ == "__main__":
    _cli()